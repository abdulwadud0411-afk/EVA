"""
Binary document extractor (.pdf, .docx).
"""
from __future__ import annotations

from pathlib import Path

from app.core.logger import get_logger
from app.knowledge.base import DocumentExtractor as BaseExtractor, ExtractedDocument, ExtractionError

logger = get_logger(__name__)


class DocumentExtractor(BaseExtractor):
    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in {".pdf", ".docx"}

    def extract(self, path: Path) -> ExtractedDocument:
        ext = path.suffix.lower()
        if ext == ".pdf":
            text = self._extract_pdf(path)
        elif ext == ".docx":
            text = self._extract_docx(path)
        else:
            raise ExtractionError(f"Unsupported extension: {ext}")

        return ExtractedDocument(
            title=path.stem,
            text=text,
            source_type="document",
            source_path=str(path),
            metadata={"extension": ext},
        )

    # ------------------------------------------------------------------ #
    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader  # type: ignore
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(
                "pypdf not installed. Run: pip install pypdf"
            ) from exc

        try:
            reader = PdfReader(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open PDF: {exc}") from exc

        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001
                continue
        return "\n".join(parts).strip()

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            import docx  # type: ignore
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(
                "python-docx not installed. Run: pip install python-docx"
            ) from exc

        try:
            doc = docx.Document(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Failed to open DOCX: {exc}") from exc

        parts = [p.text for p in doc.paragraphs if p.text]
        return "\n".join(parts).strip()